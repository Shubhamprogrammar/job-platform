import io
from pathlib import Path

import docx
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_text(file_name: str, data: bytes) -> str:
    ext = Path(file_name).suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif ext == ".txt":
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")

    text = text.strip()
    if len(text) < 50:
        raise ValueError("Could not read text from the resume. Scanned PDFs are not supported yet.")
    return text
